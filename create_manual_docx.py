import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

doc = docx.Document()

# Set standard margins (1 inch)
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    run.font.name = 'Cordia New'
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = RGBColor(27, 54, 93)
    return p

def add_subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(20)
    run = p.add_run(text)
    run.font.name = 'Cordia New'
    run.font.size = Pt(18)
    run.font.italic = True
    run.font.color.rgb = RGBColor(70, 80, 95)
    return p

def add_h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = 'Cordia New'
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(27, 54, 93)
    return p

def add_h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = 'Cordia New'
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(41, 98, 153)
    return p

def add_h3(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.name = 'Cordia New'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(50, 50, 50)
    return p

def add_body(text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.name = 'Cordia New'
        r_bold.font.size = Pt(14)
        r_bold.font.bold = True
    run = p.add_run(text)
    run.font.name = 'Cordia New'
    run.font.size = Pt(14)
    run.font.italic = italic
    return p

def add_bullet(text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.name = 'Cordia New'
        r_bold.font.size = Pt(14)
        r_bold.font.bold = True
    run = p.add_run(text)
    run.font.name = 'Cordia New'
    run.font.size = Pt(14)
    return p

def add_callout(text, title="ข้อควรระวัง / ข้อมูลเพิ่มเติม"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    
    shading = parse_xml(r'<w:shd {} w:fill="FFFBEB"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shading)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(r'''
        <w:tcBorders {} >
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="36" w:space="0" w:color="F59E0B"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    '''.format(nsdecls('w')))
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    
    r_title = p.add_run(f"📌 {title}\n")
    r_title.font.name = 'Cordia New'
    r_title.font.size = Pt(14)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(180, 83, 9)
    
    r_text = p.add_run(text)
    r_text.font.name = 'Cordia New'
    r_text.font.size = Pt(14)
    r_text.font.italic = True

def add_image_with_caption(img_path, caption):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run()
        run.add_picture(img_path, width=Inches(6.0))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(14)
        r_cap = p_cap.add_run(f"รูปที่: {caption}")
        r_cap.font.name = 'Cordia New'
        r_cap.font.size = Pt(12)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(90, 90, 90)

def style_table(table, headers, data):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        shading = parse_xml(r'<w:shd {} w:fill="1B365D"/>'.format(nsdecls('w')))
        hdr_cells[i]._tc.get_or_add_tcPr().append(shading)
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = 'Cordia New'
                r.font.size = Pt(13)
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
    
    for row_idx, row_data in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            row_cells[col_idx].text = text
            shading = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), bg_color))
            row_cells[col_idx]._tc.get_or_add_tcPr().append(shading)
            for p in row_cells[col_idx].paragraphs:
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                for r in p.runs:
                    r.font.name = 'Cordia New'
                    r.font.size = Pt(12)

img_dir = "./images"

# --- PAGE 1: TITLE ---
add_title("หนังสือคู่มือการเรียนรู้\nการสร้างเว็บพอร์ตโฟลิโอนำเสนอเกม ด้วย vibeUI และ Gemini Canvas")
add_subtitle("โครงการอบรมเชิงปฏิบัติการ เรื่อง การสร้างเกมส่งเสริมการเรียนรู้เชิงโต้ตอบด้วยปัญญาประดิษฐ์\n(Vibe Coding & Google App Script)\n\nระยะเวลาอบรม: 30 นาที | รูปแบบ: บรรยาย + ปฏิบัติควบคู่กัน")

doc.add_page_break()

# --- PAGE 2: PREFACE & CONTENTS ---
add_h1("คำนำ")
add_body("หนังสือคู่มือเล่มนี้จัดทำขึ้นเพื่อใช้ประกอบการอบรมเชิงปฏิบัติการ เรื่อง การสร้างเกมส่งเสริมการเรียนรู้เชิงโต้ตอบด้วยปัญญาประดิษฐ์ โดยมุ่งเน้นเฉพาะส่วนของการนำคลังต้นแบบ UI ที่ชื่อว่า vibeUI มาใช้งานร่วมกับเครื่องมือ Gemini Canvas เพื่อสร้างเว็บพอร์ตโฟลิโอสำหรับนำเสนอผลงานเกมที่ผู้เรียนสร้างไว้แล้ว และสามารถกดเข้าไปเล่นเกมนั้นได้จริงจากลิงก์ภายนอก")
add_body("เนื้อหาในเล่มนี้ออกแบบมาสำหรับผู้เรียนที่มีพื้นฐานคอมพิวเตอร์ทั่วไป แต่ยังไม่เคยใช้งานเครื่องมือปัญญาประดิษฐ์มาก่อน โดยอธิบายตั้งแต่พื้นฐานไปจนถึงขั้นตอนปฏิบัติจริงแบบละเอียดทีละขั้นตอน ผู้เรียนสามารถใช้หนังสือเล่มนี้อ้างอิงได้ด้วยตนเองทั้งระหว่างและหลังการอบรม")

add_h1("คำแนะนำการใช้หนังสือ")
add_body("หนังสือเล่มนี้แบ่งเนื้อหาออกเป็น 2 บทหลัก โดยแต่ละบทมีโครงสร้างเหมือนกันคือ เริ่มจากความนำ จุดประสงค์การเรียนรู้ คำศัพท์สำคัญ เนื้อหาหลัก ตัวอย่าง ขั้นตอนปฏิบัติ แบบฝึกปฏิบัติ การแก้ปัญหาเบื้องต้น และสรุปท้ายบท")
add_bullet("กล่องสีเหลือง หมายถึง ข้อมูลที่ต้องตรวจสอบเพิ่มเติมหรือข้อควรระวังพิเศษ")
add_bullet("ตาราง หมายถึง ข้อมูลเชิงเปรียบเทียบหรือคำศัพท์ที่ควรอ้างอิงได้อย่างรวดเร็ว")

add_h1("ผลลัพธ์การเรียนรู้ของหนังสือเล่มนี้")
add_body("เมื่อศึกษาหนังสือคู่มือเล่มนี้จบ ผู้เรียนสามารถ:")
add_bullet("อธิบายความหมายและโครงสร้างของคลัง vibeUI ได้")
add_bullet("เลือกหมวดหมู่และรูปแบบ UI Prompt ที่เหมาะสมกับการสร้างเว็บพอร์ตโฟลิโอนำเสนอเกมได้")
add_bullet("ปฏิบัติการสร้างเว็บพอร์ตโฟลิโอผ่าน Gemini Canvas ได้ด้วยตนเอง")
add_bullet("เชื่อมโยงปุ่มในหน้าเว็บให้ลิงก์ไปยังเกมที่เผยแพร่บนแพลตฟอร์มภายนอกได้")
add_bullet("ตรวจสอบและแก้ไขปัญหาเบื้องต้นที่พบระหว่างการใช้งานได้")

doc.add_page_break()

# --- CHAPTER 1 ---
add_h1("บทที่ 1: รู้จัก vibeUI — คลังต้นแบบ UI สำหรับสร้างงานด้วย AI")

add_h2("1. ความนำ")
add_body("ในการสร้างเกมส่งเสริมการเรียนรู้เชิงโต้ตอบด้วยปัญญาประดิษฐ์ สิ่งที่ผู้เรียนมักเจอปัญหาคือ ไม่รู้จะเริ่มออกแบบหน้าตา (UI — User Interface หรือ ส่วนติดต่อผู้ใช้งาน) ของเว็บไซต์อย่างไร และเมื่อพิมพ์คำสั่งให้ AI สร้างหน้าเว็บแบบสั้น ๆ ผลลัพธ์ที่ได้มักไม่ตรงตามที่ต้องการ เพราะ AI ไม่มีข้อมูลเพียงพอว่าเราต้องการโครงสร้างหน้าตาแบบไหน")
add_body("vibeUI คือคลังชุดคำสั่ง (Prompt)สำเร็จรูปสำหรับสร้าง UI ที่ถูกออกแบบมาให้ผู้ใช้งานสามารถเลือกโครงสร้างหน้าตาที่ต้องการ แล้วนำไปสั่งงาน AI ต่อได้ทันที โดยไม่ต้องเขียนคำอธิบายเองตั้งแต่ต้น ช่วยลดเวลาและเพิ่มความแม่นยำในการสื่อสารกับ AI")

# Real screenshot from vibeui.online
add_image_with_caption(f"{img_dir}/vibeui_real_overview.jpg", "หน้าตาคลัง Prompt สไตล์โทนสว่างของเว็บจริง vibeUI (https://vibeui.online)")

add_h2("2. จุดประสงค์การเรียนรู้")
add_bullet("อธิบายความหมายและหน้าที่ของ vibeUI ได้")
add_bullet("จำแนกหมวดหมู่ของ UI Prompt ในคลัง vibeUI ได้")
add_bullet("เลือก UI Prompt ที่เหมาะสมกับโจทย์การสร้างเว็บพอร์ตโฟลิโอนำเสนอเกมได้")
add_bullet("อธิบายหลักการทำงานของ Prompt ร่วมกับภาพอ้างอิง (Screenshot) ได้")

add_h2("3. คำศัพท์สำคัญ")
t1 = doc.add_table(rows=6, cols=3)
t1_headers = ["คำศัพท์", "ความหมาย", "ความสำคัญ"]
t1_data = [
    ["ยูไอ (UI - User Interface)", "ส่วนติดต่อผู้ใช้งานบนหน้าจอ เช่น ปุ่ม แบบฟอร์ม เมนู", "เป็นสิ่งที่ผู้ชมมองเห็นและโต้ตอบด้วยโดยตรง"],
    ["พรอมต์ (Prompt)", "ข้อความคำสั่งที่ใช้สื่อสารกับ AI เพื่อให้สร้างผลลัพธ์", "คุณภาพของ Prompt ส่งผลโดยตรงต่อคุณภาพงาน AI"],
    ["เลย์เอาต์ (Layout)", "รูปแบบการจัดวางองค์ประกอบต่างๆ บนหน้าจอ", "กำหนดตำแหน่งองค์ประกอบบนหน้าเว็บ"],
    ["สกรีนช็อต (Screenshot)", "ภาพถ่ายหน้าจอที่ใช้เป็นตัวอย่างอ้างอิงด้านสไตล์", "ช่วยให้ AI เข้าใจโทนสี ฟอนต์ และอารมณ์รวม"],
    ["แคนวาส (Canvas)", "พื้นที่ทำงานของ Gemini ที่แสดงโค้ดและตัวอย่างโต้ตอบได้", "เป็นเครื่องมือหลักที่ใช้นำ Prompt ไปสร้างงานจริง"]
]
style_table(t1, t1_headers, t1_data)

add_h2("4. แนวคิดหลัก (Core Concept)")
add_h3("4.1 vibeUI คืออะไร?")
add_body("vibeUI เป็นคลังรวมชุดคำสั่ง (Prompt) สำหรับสร้าง UI ที่จัดสรรไว้เป็นระบบ ประกอบด้วย Prompt ทั้งหมด 92 รายการ แบ่งออกเป็น 15 หมวดหมู่หลัก โดยแต่ละ Prompt จะบรรยายโครงสร้างและลักษณะของหน้าตา UI แบบหนึ่งๆ ไว้ได้อย่างชัดเจน")

add_h3("4.2 องค์ประกอบ — 15 หมวดหมู่ในคลัง vibeUI")
t2 = doc.add_table(rows=16, cols=3)
t2_headers = ["ลำดับ", "หมวดหมู่", "จำนวน Prompt / การใช้งาน"]
t2_data = [
    ["1", "Auth Forms", "6 Prompts - หน้าล็อกอิน/สมัครสมาชิก"],
    ["2", "Pricing", "8 Prompts - หน้าแสดงราคา/แพ็กเกจ"],
    ["3", "Features / Bento", "8 Prompts - หน้าแสดงจุดเด่นของฟีเจอร์"],
    ["4", "Hero Sections", "8 Prompts - ส่วนหัวเว็บไซต์ แนะนำตัวและเกม"],
    ["5", "CTA Banners", "7 Prompts - แบนเนอร์กระตุ้นการกดปุ่มกดเล่นเกม"],
    ["6", "Stats Bars", "7 Prompts - แถบแสดงตัวเลขสถิติ"],
    ["7", "Nav Bars", "8 Prompts - แถบเมนูนำทาง"],
    ["8", "Testimonials", "8 Prompts - คำรีวิว/คำชื่นชมจากผู้เล่น"],
    ["9", "Footer", "5 Prompts - ส่วนท้ายเว็บไซต์"],
    ["10", "FAQ", "5 Prompts - คำถามที่พบบ่อย"],
    ["11", "Dashboards", "6 Prompts - แผงควบคุมข้อมูล"],
    ["12", "Onboarding", "4 Prompts - ขั้นตอนแนะนำการใช้งาน"],
    ["13", "Blog / Content", "4 Prompts - หน้ารวมบทความ"],
    ["14", "Contact", "3 Prompts - หน้าติดต่อผู้พัฒนา"],
    ["15", "Bonus", "5 Prompts - หน้า 404, Loading ฯลฯ"]
]
style_table(t2, t2_headers, t2_data)

add_callout(
    "หมวดหมู่ที่มักนำมาประยุกต์ใช้กับเว็บพอร์ตโฟลิโอนำเสนอเกม ได้แก่:\n"
    "• Hero Sections (หน้าแรกแนะนำตัวและชื่อเกมพร้อมปุ่ม CTA)\n"
    "• Features/Bento (แสดงจุดเด่นของเกม ภาพหน้าจอ หรือวิธีเล่น)\n"
    "• CTA Banners (ปุ่ม 'เล่นเกมเลย' ที่ลิงก์ไป itch.io)\n"
    "• Contact/Footer (ช่องทางติดต่อผู้พัฒนา)",
    "คำแนะนำพิเศษสำหรับการนำเสนอเกม"
)

# Real screenshot of Hero Section
add_image_with_caption(f"{img_dir}/vibeui_real_hero.jpg", "ตัวอย่าง Prompt Cards ในหมวดหมู่ Hero Sections จากเว็บจริง vibeUI")

add_h2("5. วิธีการทำงาน (Input → Process → Output)")
add_bullet("Input: เลือก Prompt จากหมวดหมู่ที่ต้องการ พร้อมภาพสกรีนช็อตอ้างอิงสไตล์ (ถ้ามี)")
add_bullet("Process: นำ Prompt และภาพไปวางในเครื่องมือ AI (Gemini Canvas) จากนั้น AI จะวิเคราะห์โครงสร้างและสไตล์")
add_bullet("Output: ได้หน้า UI ที่มีโครงสร้างตาม Prompt และมีสไตล์ตามภาพอ้างอิงที่แนบไป")

doc.add_page_break()

# --- CHAPTER 2 ---
add_h1("บทที่ 2: ขั้นตอนการสร้างเว็บพอร์ตโฟลิโอด้วย vibeUI ร่วมกับ Gemini Canvas")

add_h2("1. ความนำ")
add_body("จากบทที่ 1 ผู้เรียนได้เรียนรู้วิธีเลือกหมวดหมู่และ Prompt จากคลัง vibeUI แล้ว ในบทนี้จะพาผู้เรียนไปสู่ขั้นตอนถัดไป คือการนำ Prompt ที่เลือกไว้มาสร้างเป็นเว็บพอร์ตโฟลิโอที่ใช้งานได้จริงผ่าน Gemini Canvas")

add_image_with_caption(f"{img_dir}/workflow_steps_diagram_1788958186415.jpg", "แผนภาพขั้นตอนการทำงาน 5 ขั้นตอน (5-Step Workflow)")

add_h2("2. ขั้นตอนปฏิบัติ 5 ขั้นตอนหลัก")

add_h3("ขั้นตอนที่ 1: เข้าสู่ Gemini และเปิดโหมด Canvas")
add_body("• เปิดเว็บเบราว์เซอร์ เข้าสู่ gemini.google.com ด้วยบัญชี Google\n• สังเกตช่องพิมพ์ข้อความและตัวเลือก Canvas ที่พร้อมทำงาน")

add_h3("ขั้นตอนที่ 2: เตรียม Prompt จาก vibeUI และภาพอ้างอิง")
add_body("• เปิดคลัง vibeUI เลือก Prompt หมวดหมู่ Hero Sections หรือ Bento Grid\n• คัดลอกข้อความ Prompt และเตรียมภาพหน้าจอเกมอ้างอิง")

# Light theme Canvas mockup
add_image_with_caption(f"{img_dir}/gemini_canvas_light.jpg", "พื้นที่ทำงาน Gemini Canvas โทนสว่าง แสดงช่องพิมพ์ Prompt และหน้าพรีวิวเว็บโต้ตอบได้")

add_h3("ขั้นตอนที่ 3: วางคำสั่งลงในช่องพิมพ์ของ Gemini และแนบภาพ")
add_body("• วาง (Paste) Prompt ลงในช่องพิมพ์ข้อความ\n• พิมพ์ระบุลิงก์เกมเพิ่มเติม เช่น 'และทำให้ปุ่ม Play Game เปิดลิงก์ https://mygame.itch.io ในแท็บใหม่'\n• กดปุ่มส่ง (Send) เพื่อเริ่มประมวลผล")

add_h3("ขั้นตอนที่ 4: ตรวจสอบผลลัพธ์ใน Canvas และปรับแต่งเพิ่มเติม (Refine)")
add_body("• ตรวจสอบในหน้าต่าง Preview ด้านขวาของ Canvas\n• หากต้องการแก้ไข สามารถพิมพ์คำสั่งปรับแต่งเฉพาะจุด (Iterative Refine) เช่น 'ปรับโทนสีปุ่มให้เป็นสีส้มนีออน' ได้ทันที")

# Real screenshot of Bento Section from vibeui.online
add_image_with_caption(f"{img_dir}/vibeui_real_bento.jpg", "ตัวอย่าง Prompt Cards ในหมวดหมู่ Features / Bento จากเว็บจริง vibeUI")

add_h3("ขั้นตอนที่ 5: บันทึกหรือดาวน์โหลดผลงาน")
add_body("• คลิกปุ่มเมนูด้านบนของ Canvas เลือก Export / Download ไฟล์เป็น HTML\n• สามารถนำไฟล์ .html ที่ได้ไปเปิดใช้งานบนเว็บเบราว์เซอร์ใดก็ได้")

add_h2("3. Troubleshooting & การแก้ปัญหาที่พบบ่อย")
t3 = doc.add_table(rows=4, cols=3)
t3_headers = ["ปัญหา", "สาเหตุที่เป็นไปได้", "วิธีแก้ไข"]
t3_data = [
    ["Canvas ไม่เปิดขึ้นมา", "คำสั่งไม่ได้สื่อให้ระบบสร้างโค้ดหน้าเว็บ", "ระบุคำสั่งให้ชัดเจนว่าต้องการสร้างโค้ดหน้าเว็บ HTML"],
    ["ปุ่ม 'เล่นเกม' ไม่เปิดลิงก์", "ยังไม่ได้ระบุ URL ในคำสั่ง หรือ AI ลืมเชื่อมลิงก์", "พิมพ์ Refine ระบุ URL เกมให้ชัดเจนอีกครั้ง"],
    ["สไตล์สีไม่ตรงตามต้องการ", "ไม่ได้แนบภาพอ้างอิง หรือแนบไม่สำเร็จ", "แนบภาพสกรีนช็อตตัวอย่างโทนสีใหม่อีกครั้ง"]
]
style_table(t3, t3_headers, t3_data)

add_h1("แบบฝึกปฏิบัติรวมท้ายเล่ม & Checklist")
add_body("ให้นักเรียนสร้างหน้าเว็บพอร์ตโฟลิโอ 1 หน้า สำหรับนำเสนอเกมของตนเอง โดยปฏิบัติตาม Checklist ดังนี้:")
add_bullet("[  ] กำหนดชื่อเกมและเตรียมลิงก์เกมเรียบร้อย")
add_bullet("[  ] เลือก Prompt จากคลัง vibeUI ที่เหมาะสม")
add_bullet("[  ] วาง Prompt ใน Gemini Canvas และเพิ่มลิงก์เกม")
add_bullet("[  ] ตรวจสอบ Preview และทดสอบการกดปุ่มไปเล่นเกมจริง")
add_bullet("[  ] ดาวน์โหลดไฟล์ HTML เก็บไว้ใช้งาน")

doc.save("คู่มือการเรียนรู้_vibeUI_GeminiCanvas.docx")
print("Successfully generated Word document with real light-theme images!")
